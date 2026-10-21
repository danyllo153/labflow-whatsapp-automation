"""Planilha Excel (.xlsx) do que está filtrado numa página.

Spec: 07-app-pwa.md (v1.0), "5." (exportar o filtrado); 08-fluxos.md A11. Tarefa: APP-009.
Formal e sem emoji: aba com os dados e aba "Filtros" com os filtros usados e a data da exportação.
"""

from datetime import date, datetime
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo

from fastapi import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_FUSO = ZoneInfo("America/Sao_Paulo")
_CABECALHO = PatternFill("solid", fgColor="F6B04A")  # laranja do app (07, Visual)


def _celula(valor: Any) -> Any:
    if isinstance(valor, datetime):  # Excel não guarda fuso: mostra a hora de São Paulo
        return valor.astimezone(_FUSO).replace(tzinfo=None) if valor.tzinfo else valor
    return valor


def planilha(titulo: str, colunas: list[tuple[str, str]], linhas: list[dict[str, Any]],
             filtros: list[tuple[str, Any]]) -> bytes:
    livro = Workbook()
    aba = livro.active
    aba.title = titulo
    aba.append([rotulo for _, rotulo in colunas])
    for celula in aba[1]:
        celula.font = Font(bold=True)
        celula.fill = _CABECALHO
    for linha in linhas:
        aba.append([_celula(linha.get(chave)) for chave, _ in colunas])
    for i, (chave, rotulo) in enumerate(colunas, start=1):
        maior = max([len(rotulo)] + [len(str(l.get(chave) or "")) for l in linhas[:500]])
        aba.column_dimensions[get_column_letter(i)].width = min(maior + 2, 40)
        formato = "dd/mm/yyyy hh:mm" if any(isinstance(l.get(chave), datetime) for l in linhas[:50]) else (
            "dd/mm/yyyy" if any(isinstance(l.get(chave), date) for l in linhas[:50]) else None)
        if formato:
            for celula in aba[get_column_letter(i)][1:]:
                celula.number_format = formato
    aba.freeze_panes = "A2"
    aba.auto_filter.ref = aba.dimensions

    info = livro.create_sheet("Filtros")
    info.append(["LabFlow", titulo])
    info.append(["Exportado em", datetime.now(_FUSO).strftime("%d/%m/%Y %H:%M")])
    info.append(["Linhas", len(linhas)])
    info.append([])
    info.append(["Filtro", "Valor"])
    for nome, valor in filtros:
        info.append([nome, valor.strftime("%d/%m/%Y") if isinstance(valor, date) else (valor if valor not in (None, "") else "(todos)")])
    info.column_dimensions["A"].width = 20
    info.column_dimensions["B"].width = 30

    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()


def resposta_xlsx(conteudo: bytes, nome: str) -> Response:
    hoje = datetime.now(_FUSO).strftime("%Y-%m-%d")
    return Response(
        content=conteudo,
        media_type=MIME_XLSX,
        headers={"Content-Disposition": f'attachment; filename="labflow-{nome}-{hoje}.xlsx"'},
    )
