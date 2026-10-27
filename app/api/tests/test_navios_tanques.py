"""APP-019b: Tanques terra e Navios com status, e a página de cada um (pedido do Danyllo em 23/10/2026).

Status: Concluído = nada pendente (leituras, drops, desvios, C.T/B.L do embarque, TAB/Coliformes/Howard);
Atrasado = algo passou do prazo; Em andamento = o resto. Conferido contra as views do Power BI.
"""

from urllib.parse import quote

import pytest

SENHA = "Senha-de-teste-1"


@pytest.fixture
def logado(cliente, criar_pessoa):
    def _logar(cargo="Analista"):
        criar_pessoa(usuario=f"{cargo.lower()}.teste", cargo=cargo)
        assert cliente.post("/login", json={"usuario": f"{cargo.lower()}.teste", "senha": SENHA}).status_code == 200

    return _logar


def _um(conexao, sql, *p):
    return conexao.execute(sql, p).fetchone()[0]


def _status_esperado(atrasados, pendentes):
    return "Atrasado" if atrasados else ("Concluído" if pendentes == 0 else "Em andamento")


# ---------------------------------------------------------------- Tanques terra
def test_tanques_terra_um_por_coleta_de_terra(cliente, logado, conexao):
    logado()
    corpo = cliente.get("/paginas/tanques-terra", params={"por_pagina": 200}).json()
    assert corpo["total"] == _um(conexao, "SELECT count(*) FROM bi.coletas WHERE origem = 'Terra'") > 0
    assert all(l["origem"] == "Terra" for l in corpo["linhas"])


def test_status_de_cada_tanque_segue_a_regra(cliente, logado, conexao):
    logado()
    for l in cliente.get("/paginas/tanques-terra", params={"por_pagina": 200}).json()["linhas"]:
        cid, tq, dia = l["coleta_id"], l["tanque"], l["data_coleta"]
        atras = (
            _um(conexao, "SELECT count(*) FROM bi.leituras_nfc WHERE coleta_id = %s AND situacao = 'Atrasado'", cid)
            + _um(conexao, "SELECT count(*) FROM bi.drops WHERE coleta_id = %s AND situacao = 'Atrasado'", cid)
            + _um(conexao, "SELECT count(*) FROM bi.desvios WHERE origem = 'Terra' AND tanque = %s AND data_coleta = %s AND situacao = 'Atrasado'", tq, dia)
            + _um(conexao, "SELECT count(*) FROM bi.testes WHERE origem = 'Tank farm' AND tanque = %s AND data_coleta = %s AND situacao = 'Atrasado'", tq, dia)
        )
        pend = (
            _um(conexao, "SELECT count(*) FROM bi.leituras_nfc WHERE coleta_id = %s AND NOT lida", cid)
            + _um(conexao, "SELECT count(*) FROM bi.drops WHERE coleta_id = %s AND NOT concluido", cid)
            + _um(conexao, "SELECT count(*) FROM bi.desvios WHERE origem = 'Terra' AND tanque = %s AND data_coleta = %s AND status <> 'Concluído'", tq, dia)
            + _um(conexao, "SELECT count(*) FROM bi.testes WHERE origem = 'Tank farm' AND tanque = %s AND data_coleta = %s AND situacao <> 'Concluído'", tq, dia)
        )
        assert l["status"] == _status_esperado(atras, pend), (tq, dia)


def test_tanques_terra_filtro_de_status(cliente, logado):
    logado()
    todos = cliente.get("/paginas/tanques-terra", params={"por_pagina": 200}).json()["linhas"]
    for st in ("Concluído", "Atrasado", "Em andamento"):
        n = cliente.get("/paginas/tanques-terra", params={"status": st}).json()["total"]
        assert n == sum(l["status"] == st for l in todos)


def test_pagina_do_tanque_traz_leituras_com_metodo_drops_e_tab(cliente, logado, conexao):
    logado()
    cid = _um(conexao, "SELECT coleta_id FROM bi.leituras_nfc l JOIN bi.coletas c USING (coleta_id) WHERE c.origem = 'Terra' AND l.analise = 'Psicrotróficos' LIMIT 1")
    d = cliente.get(f"/paginas/tanques-terra/{cid}").json()
    assert len(d["leituras"]) == _um(conexao, "SELECT count(*) FROM bi.leituras_nfc WHERE coleta_id = %s", cid)
    metodos = {l["metodo"] for l in d["leituras"] if l["analise"] == "Psicrotróficos"}
    assert metodos == {"Profundidade", "Superficie"}  # as linhas que pareciam repetidas são dois métodos
    assert len(d["drops"]) == _um(conexao, "SELECT count(*) FROM bi.drops WHERE coleta_id = %s", cid)
    assert d["resumo"]["coleta_id"] == cid


def test_tanque_que_nao_existe_da_404(cliente, logado):
    logado()
    assert cliente.get("/paginas/tanques-terra/999999").status_code == 404


# ---------------------------------------------------------------- Navios
def test_navios_traz_todos_nfc_e_fcoj(cliente, logado, conexao):
    logado()
    linhas = cliente.get("/paginas/navios", params={"por_pagina": 200}).json()["linhas"]
    esperados = {
        r[0]
        for r in conexao.execute(
            "SELECT navio FROM bi.coletas WHERE origem = 'Navio' UNION SELECT navio FROM bi.contagens WHERE origem = 'Embarque'"
        ).fetchall()
    }
    assert {l["navio"] for l in linhas} == esperados
    assert {l["produto"] for l in linhas} >= {"NFC", "FCOJ"}
    status = [l["status"] for l in linhas]
    assert status == sorted(status, key=lambda s: s == "Concluído")  # os que faltam alguma coisa primeiro


def test_navio_fcoj_soma_amostras_e_tab_do_embarque(cliente, logado, conexao):
    logado()
    nome = _um(conexao, "SELECT navio FROM bi.contagens WHERE origem = 'Embarque' ORDER BY navio LIMIT 1")
    d = cliente.get(f"/paginas/navios/{quote(nome)}").json()
    assert d["resumo"]["navio"] == nome
    assert d["resumo"]["tanques_nfc"] == len({t["tanque"] for t in d["tanques"]})
    pend = sum(e["pendentes"] for e in d["embarques"])
    assert pend == _um(conexao, "SELECT count(*) FROM bi.contagens WHERE origem = 'Embarque' AND navio = %s AND NOT lido", nome)
    for e in d["embarques"]:
        abertos = _um(
            conexao,
            "SELECT count(*) FROM bi.testes WHERE origem = 'Embarque' AND embarque = %s AND situacao <> 'Concluído'",
            f"{nome} linha {e['linha']} fase {e['fase']}",
        )
        assert e["testes_abertos"] == abertos
        atras = e["atrasadas"] + _um(
            conexao, "SELECT count(*) FROM bi.testes WHERE origem = 'Embarque' AND embarque = %s AND situacao = 'Atrasado'",
            f"{nome} linha {e['linha']} fase {e['fase']}",
        )
        assert e["status"] == _status_esperado(atras, e["pendentes"] + abertos)


def test_navio_traz_cada_amostra_do_fcoj_com_o_ct_e_o_bl(cliente, logado, conexao):
    """Pedido de 27/10: tocar no navio mostra todas as análises do FCOJ embarcado, com os resultados."""
    logado()
    nome = _um(conexao, "SELECT navio FROM bi.contagens WHERE origem = 'Embarque' ORDER BY navio LIMIT 1")
    amostras = cliente.get(f"/paginas/navios/{quote(nome)}").json()["amostras"]
    assert len(amostras) == _um(
        conexao,
        "SELECT count(DISTINCT (linha, fase, load, lote_amostra)) FROM bi.contagens WHERE origem = 'Embarque' AND navio = %s",
        nome,
    )
    a = amostras[0]
    assert {"linha", "fase", "load", "amostra", "ct", "bl72", "bl120", "situacao"} <= set(a)
    ct = _um(
        conexao,
        "SELECT max(resultado) FROM bi.contagens WHERE origem = 'Embarque' AND navio = %s AND linha = %s AND fase = %s "
        "AND load IS NOT DISTINCT FROM %s AND lote_amostra = %s AND analise = 'C.T 48h'",
        nome, a["linha"], a["fase"], a["load"], a["amostra"],
    )
    assert a["ct"] == ct


def test_pagina_do_load_traz_os_lotes_e_as_compostas(cliente, logado, conexao):
    logado()
    load = _um(conexao, "SELECT load FROM bi.contagens WHERE origem = 'Recebimento' ORDER BY load LIMIT 1")
    d = cliente.get(f"/paginas/loads/{quote(load)}").json()
    assert d["resumo"]["load"] == load
    assert len(d["lotes"]) == _um(
        conexao, "SELECT count(DISTINCT (item, numero)) FROM bi.contagens WHERE origem = 'Recebimento' AND load = %s", load
    )
    assert len(d["testes"]) == _um(conexao, "SELECT count(*) FROM bi.testes WHERE origem = 'Recebimento' AND load_ou_tanque = %s", load)
    assert [l["lote"] for l in d["lotes"]] == sorted(l["lote"] for l in d["lotes"])
    assert cliente.get("/paginas/loads/NAO-EXISTE").status_code == 404


def test_ultimos_navios_embarcados_do_mais_novo_ao_mais_antigo(cliente, logado, conexao):
    logado()
    d = cliente.get("/analises/embarque/navios", params={"por_pagina": 2}).json()
    assert d["total"] == _um(conexao, "SELECT count(DISTINCT navio) FROM bi.contagens WHERE origem = 'Embarque' AND navio IS NOT NULL")
    datas = [n["ultima_data"] for n in d["navios"]]
    assert datas == sorted(datas, reverse=True) and len(datas) <= 2
    if d["total"] > 2:
        seguinte = cliente.get("/analises/embarque/navios", params={"por_pagina": 2, "pagina": 2}).json()["navios"]
        assert seguinte and seguinte[0]["navio"] not in {n["navio"] for n in d["navios"]}


@pytest.mark.parametrize("origem", ["terra", "navio"])
def test_coletas_dos_tanques_por_dia_e_pessoa(cliente, logado, conexao, origem):
    logado()
    coletas = cliente.get("/analises/tanques/coletas", params={"origem": origem}).json()
    total = _um(
        conexao,
        "SELECT count(*) FROM bi.coletas WHERE origem = %s AND data_coleta > current_date - 366",
        "Terra" if origem == "terra" else "Navio",
    )
    assert sum(c["n"] for c in coletas) <= total  # o "hoje" da API é o do relógio do laboratório (pode ser um dia antes)
    assert all(c["pessoa"] and c["alvo"] for c in coletas)  # o tanque (terra) ou o navio: as barras empilham por ele


def test_ultimos_navios_dos_tanques(cliente, logado, conexao):
    """27/10: a página dos Tanques navio mostra os últimos navios com tanque coletado, como o rastreio do embarque."""
    logado()
    d = cliente.get("/analises/tanques/navios", params={"por_pagina": 3}).json()
    assert d["total"] == _um(conexao, "SELECT count(DISTINCT navio) FROM bi.leituras_nfc WHERE origem = 'Navio' AND navio IS NOT NULL")
    datas = [n["ultima_data"] for n in d["navios"]]
    assert datas == sorted(datas, reverse=True) and all(n["amostras"] >= 1 for n in d["navios"])


def test_navio_nfc_traz_os_tanques_do_navio(cliente, logado, conexao):
    logado()
    nome = _um(conexao, "SELECT navio FROM bi.coletas WHERE origem = 'Navio' LIMIT 1")
    d = cliente.get(f"/paginas/navios/{quote(nome)}").json()
    assert len(d["tanques"]) == _um(conexao, "SELECT count(*) FROM bi.coletas WHERE origem = 'Navio' AND navio = %s", nome)
    assert all(t["navio"] == nome and t["origem"] == "Navio" for t in d["tanques"])
    # o tanque do navio abre na mesma página de tanque
    assert cliente.get(f"/paginas/tanques-terra/{d['tanques'][0]['coleta_id']}").status_code == 200


def test_navios_filtros_de_produto_e_status(cliente, logado):
    logado()
    so_fcoj = cliente.get("/paginas/navios", params={"produto": "FCOJ"}).json()["linhas"]
    assert so_fcoj and all("FCOJ" in l["produto"] for l in so_fcoj)
    concluidos = cliente.get("/paginas/navios", params={"status": "Concluído"}).json()["linhas"]
    assert all(l["status"] == "Concluído" for l in concluidos)


def test_navio_que_nao_existe_da_404(cliente, logado):
    logado()
    assert cliente.get("/paginas/navios/NAO%20EXISTE%20999").status_code == 404


@pytest.mark.parametrize("caminho", ["/paginas/tanques-terra", "/paginas/navios", "/paginas/navios/D.SKY%20123", "/paginas/tanques-terra/1",
                                     "/paginas/loads/90001"])
def test_coletor_nao_ve(cliente, logado, caminho):
    logado("Coletor")
    assert cliente.get(caminho).status_code == 403


@pytest.mark.parametrize("pagina", ["tanques-terra", "navios"])
def test_excel_igual_a_pagina(cliente, logado, pagina):
    from io import BytesIO

    from openpyxl import load_workbook

    logado()
    total = cliente.get(f"/paginas/{pagina}").json()["total"]
    folha = load_workbook(BytesIO(cliente.get(f"/paginas/{pagina}/excel").content)).worksheets[0]
    assert folha.max_row - 1 == total > 0
