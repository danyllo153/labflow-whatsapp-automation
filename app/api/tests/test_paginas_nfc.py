"""APP-017: páginas Tanques, Drops e Desvios (07 v1.0, "5. Páginas do laboratório"; 08 A11).

Os números são conferidos contra as próprias views do Power BI (bi.leituras_nfc, bi.drops, bi.desvios) no banco de teste.
"""

from io import BytesIO

import pytest
from openpyxl import load_workbook

from labflow_api.excel import MIME_XLSX

SENHA = "Senha-de-teste-1"


@pytest.fixture
def logado(cliente, criar_pessoa):
    def _logar(cargo="Analista"):
        criar_pessoa(usuario=f"{cargo.lower()}.teste", cargo=cargo)
        assert cliente.post("/login", json={"usuario": f"{cargo.lower()}.teste", "senha": SENHA}).status_code == 200

    return _logar


def _conta(conexao, sql, *params):
    return conexao.execute(sql, params).fetchone()[0]


# (página, view, filtros da API, WHERE equivalente na view)
CASOS = [
    ("tanques", "bi.leituras_nfc", {}, "true"),
    ("tanques", "bi.leituras_nfc", {"origem": "Terra", "analise": "B.L"}, "origem = 'Terra' AND analise = 'B.L'"),
    ("tanques", "bi.leituras_nfc", {"resultado": "Não ok"}, "resultado = 'Não ok'"),
    ("tanques", "bi.leituras_nfc", {"frasco": "Stress", "etapa": "Leitura final"}, "frasco = 'Stress' AND etapa = 'Leitura final'"),
    ("tanques", "bi.leituras_nfc", {"situacao": "Atrasado"}, "situacao = 'Atrasado'"),
    ("tanques", "bi.leituras_nfc", {"data_de": "2026-10-10", "data_ate": "2026-10-15"}, "data_coleta BETWEEN '2026-10-10' AND '2026-10-15'"),
    ("drops", "bi.drops", {}, "true"),
    ("drops", "bi.drops", {"drop": "D5", "situacao": "Concluído"}, "drop = 'D5' AND situacao = 'Concluído'"),
    ("drops", "bi.drops", {"origem": "Navio"}, "origem = 'Navio'"),
    ("desvios", "bi.desvios", {}, "true"),
    ("desvios", "bi.desvios", {"aberto": "Abertos"}, "status <> 'Concluído'"),
    ("desvios", "bi.desvios", {"aberto": "Concluídos", "tipo": "Drop"}, "status = 'Concluído' AND tipo = 'Drop'"),
    ("desvios", "bi.desvios", {"tipo": "Tanque"}, "tipo = 'Tanque'"),
]


@pytest.mark.parametrize(("pagina", "view", "filtros", "onde"), CASOS)
def test_totais_batem_com_as_views_do_power_bi(cliente, logado, conexao, pagina, view, filtros, onde):
    logado()
    r = cliente.get(f"/paginas/{pagina}", params={**filtros, "por_pagina": 200})
    assert r.status_code == 200, r.text
    corpo = r.json()
    assert corpo["total"] == _conta(conexao, f"SELECT count(*) FROM {view} WHERE {onde}")
    assert len(corpo["linhas"]) == min(corpo["total"], 200)


def test_o_banco_de_teste_tem_dados_nas_tres_paginas(cliente, logado):
    logado()
    for pagina in ("tanques", "drops", "desvios"):
        assert cliente.get(f"/paginas/{pagina}").json()["total"] > 0, pagina


def test_tanque_e_navio_filtram_como_nas_coletas(cliente, logado, conexao):
    logado()
    tanque, navio = conexao.execute(
        "SELECT tanque, navio FROM bi.leituras_nfc WHERE origem = 'Navio' LIMIT 1"
    ).fetchone()
    r = cliente.get("/paginas/tanques", params={"tanque": tanque.lower(), "navio": navio.split()[0].lower()})
    linhas = r.json()["linhas"]
    assert linhas and all(l["tanque"] == tanque and navio.split()[0] in l["navio"] for l in linhas)


def test_desvios_abertos_vem_primeiro(cliente, logado):
    logado()
    status = [l["status"] for l in cliente.get("/paginas/desvios", params={"por_pagina": 200}).json()["linhas"]]
    assert "Concluído" in status and any(s != "Concluído" for s in status)  # o banco de teste tem os dois
    assert status == sorted(status, key=lambda s: s == "Concluído")  # todos os abertos antes dos concluídos


def test_quem_nos_desvios_procura_quem_abriu_ou_deu_o_resultado(cliente, logado, conexao):
    logado()
    nome = _conta(conexao, "SELECT aberto_por FROM bi.desvios WHERE aberto_por IS NOT NULL LIMIT 1")
    esperado = _conta(
        conexao, "SELECT count(*) FROM bi.desvios WHERE aberto_por ILIKE %s OR resultado_por ILIKE %s", f"%{nome}%", f"%{nome}%"
    )
    assert cliente.get("/paginas/desvios", params={"quem": nome}).json()["total"] == esperado


@pytest.mark.parametrize(("pagina", "filtro"), [("tanques", {"analise": "WORT"}), ("drops", {"drop": "D7"}), ("desvios", {"aberto": "talvez"})])
def test_filtro_fora_da_lista_da_422(cliente, logado, pagina, filtro):
    logado()
    assert cliente.get(f"/paginas/{pagina}", params=filtro).status_code == 422


@pytest.mark.parametrize("pagina", ["tanques", "drops", "desvios"])
def test_coletor_nao_ve_as_paginas_de_analise(cliente, logado, pagina):
    logado("Coletor")
    assert cliente.get(f"/paginas/{pagina}").status_code == 403
    assert cliente.get(f"/paginas/{pagina}/excel").status_code == 403


@pytest.mark.parametrize("cargo", ["Admin", "Consultor"])
def test_outros_cargos_veem(cliente, logado, cargo):
    logado(cargo)
    assert cliente.get("/paginas/drops").status_code == 200


@pytest.mark.parametrize("pagina", ["tanques", "drops", "desvios"])
def test_sem_login_da_401(cliente, pagina):
    assert cliente.get(f"/paginas/{pagina}").status_code == 401


@pytest.mark.parametrize(
    ("pagina", "filtros", "aba"),
    [("tanques", {"resultado": "Não ok"}, "Tanques"), ("drops", {"drop": "D10"}, "Drops"), ("desvios", {"aberto": "Abertos"}, "Desvios")],
)
def test_excel_tem_as_mesmas_linhas_da_pagina_filtrada(cliente, logado, pagina, filtros, aba):
    logado()
    total = cliente.get(f"/paginas/{pagina}", params=filtros).json()["total"]
    r = cliente.get(f"/paginas/{pagina}/excel", params=filtros)
    assert r.status_code == 200 and r.headers["content-type"] == MIME_XLSX
    livro = load_workbook(BytesIO(r.content))
    assert livro.worksheets[0].title == aba
    assert livro.worksheets[0].max_row - 1 == total
