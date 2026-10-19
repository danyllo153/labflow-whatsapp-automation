"""APP-009: exportar para Excel o que está filtrado (07 v1.0, "5."; 08 A11; decisão de 15/10: livre, sem registro)."""

from io import BytesIO

import pytest
from openpyxl import load_workbook

from labflow_api.excel import MIME_XLSX

SENHA = "Senha-de-teste-1"


@pytest.fixture
def logado(cliente, criar_pessoa):
    def _logar(cargo="Operador"):
        uid = criar_pessoa(usuario=f"{cargo.lower()}.teste", cargo=cargo)
        assert cliente.post("/login", json={"usuario": f"{cargo.lower()}.teste", "senha": SENHA}).status_code == 200
        return uid

    return _logar


def _abrir(resposta):
    assert resposta.status_code == 200, resposta.text
    assert resposta.headers["content-type"] == MIME_XLSX
    assert resposta.headers["content-disposition"].startswith('attachment; filename="labflow-')
    return load_workbook(BytesIO(resposta.content))


def test_excel_de_coletas_tem_as_mesmas_linhas_da_pagina_filtrada(cliente, logado):
    logado()
    filtros = {"origem": "Navio"}
    total = cliente.get("/paginas/coletas", params=filtros).json()["total"]
    livro = _abrir(cliente.get("/paginas/coletas/excel", params=filtros))
    aba = livro["Coletas"]
    assert [c.value for c in aba[1]] == [
        "Data da coleta", "Origem", "Tanque", "Navio", "Coletado por", "Digitado por", "Registrado em"
    ]
    assert aba.max_row - 1 == total > 0
    assert {aba.cell(row=r, column=2).value for r in range(2, aba.max_row + 1)} == {"Navio"}
    assert aba.freeze_panes == "A2"


def test_aba_filtros_mostra_o_que_foi_usado(cliente, logado):
    logado()
    livro = _abrir(cliente.get("/paginas/coletas/excel", params={"tanque": "80", "data_de": "2026-01-01"}))
    filtros = {linha[0]: linha[1] for linha in livro["Filtros"].iter_rows(min_row=5, values_only=True) if linha[0]}
    assert filtros["Tanque"] == "80"
    assert filtros["Data de"] == "01/01/2026"
    assert filtros["Navio"] == "(todos)"


def test_excel_de_recebimento_igual_a_pagina(cliente, logado):
    logado("Consultor")  # Consultor também exporta (decisão de 15/10)
    filtros = {"situacao": "Não ok"}
    total = cliente.get("/paginas/recebimento", params=filtros).json()["total"]
    aba = _abrir(cliente.get("/paginas/recebimento/excel", params=filtros))["Recebimento"]
    assert aba.max_row - 1 == total > 0
    coluna_situacao = [c.value for c in aba[1]].index("Situação") + 1
    assert {aba.cell(row=r, column=coluna_situacao).value for r in range(2, aba.max_row + 1)} == {"Não ok"}


def test_datas_e_horas_sem_fuso_no_excel(cliente, logado):
    logado()
    aba = _abrir(cliente.get("/paginas/coletas/excel"))["Coletas"]
    hora = aba.cell(row=2, column=7).value
    assert hora is not None and hora.tzinfo is None  # Excel não aceita fuso; vai a hora de São Paulo


def test_coletor_exporta_so_as_proprias_coletas(cliente, logado, conexao):
    uid = logado("Coletor")
    conexao.execute(
        "INSERT INTO coletas (origem, tanque, data_coleta, registrado_por) VALUES ('terra', '92', CURRENT_DATE, %s)",
        (uid,),
    )
    aba = _abrir(cliente.get("/paginas/coletas/excel"))["Coletas"]
    assert aba.max_row == 2 and aba.cell(row=2, column=3).value == "92"
    assert cliente.get("/paginas/recebimento/excel").status_code == 403


def test_sem_login_nao_exporta(cliente):
    assert cliente.get("/paginas/coletas/excel").status_code == 401
