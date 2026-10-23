"""APP-018: páginas Embarque, TAB e Coliformes, Recebimento NFC e Visão geral (Navios: test_navios_tanques.py) (07 v1.0, §5; 08 A11).

Os números são conferidos contra as próprias views do Power BI no banco de teste.
"""

from io import BytesIO

import pytest
from openpyxl import load_workbook

from labflow_api.excel import MIME_XLSX

SENHA = "Senha-de-teste-1"


@pytest.fixture
def logado(cliente, criar_pessoa):
    def _logar(cargo="Analista"):
        criar_pessoa(usuario=f"{cargo.lower()}.teste", cargo=cargo)
        assert cliente.post("/login", json={"usuario": f"{cargo.lower()}.teste", "senha": SENHA}).status_code == 200

    return _logar


@pytest.fixture
def recebimento_nfc(conexao):
    """O banco de teste não tem NFC recebido: cria 2 lotes fictícios (fresh e past), um com C.T lido. Desfeito no fim."""
    load_id = conexao.execute("SELECT id FROM loads ORDER BY id LIMIT 1").fetchone()[0]
    ids = [
        conexao.execute(
            "INSERT INTO recebimento_nfc (load_id, lote, amostra, diluicao, data_recebimento) "
            "VALUES (%s, %s, %s, %s, CURRENT_DATE - 6) RETURNING id",
            (load_id, lote, amostra, dil),
        ).fetchone()[0]
        for lote, amostra, dil in [(1, "fresh", 3), (1, "past", 2)]
    ]
    conexao.execute(
        "INSERT INTO contagens_nfc (analise, recebimento_nfc_id, sinal, valor, valor_lido, diluicao) "
        "VALUES ('CT', %s, '=', 14000, 14, 3)",
        (ids[0],),
    )
    return ids


def _conta(conexao, sql, *params):
    return conexao.execute(sql, params).fetchone()[0]


CASOS = [
    ("embarque", "SELECT count(*) FROM bi.contagens WHERE origem = 'Embarque'", {}),
    ("embarque", "SELECT count(*) FROM bi.contagens WHERE origem = 'Embarque' AND analise = 'C.T 48h'", {"analise": "C.T 48h"}),
    ("embarque", "SELECT count(*) FROM bi.contagens WHERE origem = 'Embarque' AND lido AND alarme", {"situacao": "Não ok"}),
    ("embarque", "SELECT count(*) FROM bi.contagens WHERE origem = 'Embarque' AND NOT lido AND situacao = 'Atrasado'", {"situacao": "Atrasado"}),
    ("testes", "SELECT count(*) FROM bi.testes", {}),
    ("testes", "SELECT count(*) FROM bi.testes WHERE teste = 'TAB' AND origem = 'Recebimento'", {"teste": "TAB", "origem": "Recebimento"}),
    ("testes", "SELECT count(*) FROM bi.testes WHERE resultado = 'Positivo'", {"resultado": "Positivo"}),
    ("testes", "SELECT count(*) FROM bi.testes WHERE teste = 'Howard'", {"teste": "Howard"}),
    ("testes", "SELECT count(*) FROM bi.testes WHERE status = 'No caldo'", {"status": "No caldo"}),
]


@pytest.mark.parametrize(("pagina", "sql", "filtros"), CASOS)
def test_totais_batem_com_as_views_do_power_bi(cliente, logado, conexao, pagina, sql, filtros):
    logado()
    r = cliente.get(f"/paginas/{pagina}", params={**filtros, "por_pagina": 200})
    assert r.status_code == 200, r.text
    assert r.json()["total"] == _conta(conexao, sql)


def test_testes_abertos_vem_antes_dos_concluidos(cliente, logado):
    logado()
    sit = [l["situacao"] for l in cliente.get("/paginas/testes", params={"por_pagina": 200}).json()["linhas"]]
    assert sit == sorted(sit, key=lambda s: s == "Concluído") and "Concluído" in sit


def test_referencia_procura_load_ou_tanque(cliente, logado, conexao):
    logado()
    esperado = _conta(conexao, "SELECT count(*) FROM bi.testes WHERE load_ou_tanque ILIKE %s", "%Tanque 8%")
    assert cliente.get("/paginas/testes", params={"referencia": "tanque 8"}).json()["total"] == esperado > 0


def test_recebimento_nfc_mostra_valor_lido_e_com_diluicao(cliente, logado, recebimento_nfc):
    logado()
    corpo = cliente.get("/paginas/recebimento-nfc", params={"por_pagina": 200}).json()
    assert corpo["total"] == 6  # 2 lotes x (C.T 48h, B.L 72h, B.L 120h)
    lido = next(l for l in corpo["linhas"] if l["situacao"] == "Lido")
    assert (lido["analise"], lido["amostra"], lido["diluicao"], lido["valor_lido"], lido["resultado"]) == (
        "C.T 48h", "Fresh", "10^3", "14", "14000"
    )  # números decimais do banco vão como texto, igual às outras páginas
    assert "alarme" not in lido and lido["situacao"] not in ("Ok", "Não ok")  # NFC recebido não usa limites (027)
    assert cliente.get("/paginas/recebimento-nfc", params={"amostra": "Pasteurizado"}).json()["total"] == 3
    assert cliente.get("/paginas/recebimento-nfc", params={"situacao": "Lido"}).json()["total"] == 1


# ---------------------------------------------------------------- Visão geral
def test_visao_geral_padrao_sao_30_dias_ate_hoje(cliente, logado, conexao):
    logado()
    corpo = cliente.get("/paginas/visao-geral").json()
    hoje = _conta(conexao, "SELECT CURRENT_DATE")
    assert corpo["data_ate"] == hoje.isoformat()
    assert corpo["data_de"] == _conta(conexao, "SELECT CURRENT_DATE - 29").isoformat()


def test_visao_geral_numeros_batem_com_as_views(cliente, logado, conexao):
    logado()
    de, ate = "2026-10-01", "2026-10-31"
    corpo = cliente.get("/paginas/visao-geral", params={"data_de": de, "data_ate": ate}).json()
    n = {x["chave"]: x["valor"] for x in corpo["numeros"]}
    assert n["coletas"] == _conta(conexao, "SELECT count(*) FROM bi.coletas WHERE data_coleta BETWEEN %s AND %s", de, ate)
    assert n["recebimento_nao_ok"] == _conta(
        conexao, "SELECT count(*) FROM bi.contagens WHERE origem = 'Recebimento' AND lido AND alarme AND prevista BETWEEN %s AND %s", de, ate
    )
    assert n["testes_positivos"] == _conta(
        conexao, "SELECT count(*) FROM bi.testes WHERE resultado = 'Positivo' AND data_feito BETWEEN %s AND %s", de, ate
    )
    assert n["desvios_abertos"] == _conta(conexao, "SELECT count(*) FROM bi.desvios WHERE status <> 'Concluído'")


def test_visao_geral_pendencias_batem_com_o_inicio(cliente, logado):
    logado()
    pend = cliente.get("/pendencias-hoje").json()
    blocos = cliente.get("/paginas/visao-geral").json()["pendencias"]
    assert [b["bloco"] for b in blocos][0] == "FCOJ Recebimento" and len(blocos) == 6
    assert sum(b["vence_hoje"] for b in blocos) == pend["vence_hoje"]
    assert sum(b["atrasados"] for b in blocos) == pend["atrasados"]


def test_visao_geral_alertas_sao_os_nao_ok_mais_recentes(cliente, logado):
    logado()
    alertas = cliente.get("/paginas/visao-geral", params={"data_de": "2026-01-01"}).json()["alertas"]
    assert 0 < len(alertas) <= 20
    assert {a["pagina"] for a in alertas} <= {"recebimento", "embarque", "tanques"}
    assert all(a["caminho"].startswith(("/paginas/recebimento?load=", "/paginas/navios/", "/paginas/tanques-terra/")) for a in alertas)
    datas = [a["data"] for a in alertas if a["data"]]
    assert datas == sorted(datas, reverse=True)


def test_visao_geral_periodo_invertido_da_422(cliente, logado):
    logado()
    assert cliente.get("/paginas/visao-geral", params={"data_de": "2026-10-10", "data_ate": "2026-10-01"}).status_code == 422


# ---------------------------------------------------------------- Permissões, filtros e Excel
PAGINAS = ["embarque", "testes", "recebimento-nfc"]


@pytest.mark.parametrize("pagina", [*PAGINAS, "visao-geral"])
def test_coletor_nao_ve(cliente, logado, pagina):
    logado("Coletor")
    assert cliente.get(f"/paginas/{pagina}").status_code == 403


@pytest.mark.parametrize("pagina", [*PAGINAS, "visao-geral"])
def test_sem_login_da_401(cliente, pagina):
    assert cliente.get(f"/paginas/{pagina}").status_code == 401


@pytest.mark.parametrize(
    ("pagina", "filtro"),
    [("embarque", {"analise": "B.L 48h"}), ("testes", {"teste": "Salmonela"}), ("recebimento-nfc", {"amostra": "Past"})],
)
def test_filtro_fora_da_lista_da_422(cliente, logado, pagina, filtro):
    logado()
    assert cliente.get(f"/paginas/{pagina}", params=filtro).status_code == 422


@pytest.mark.parametrize(
    ("pagina", "filtros", "aba"),
    [
        ("embarque", {"situacao": "Atrasado"}, "Embarque"),
        ("testes", {"teste": "Coliformes"}, "TAB e Coliformes"),
        ("recebimento-nfc", {}, "Recebimento NFC"),
    ],
)
def test_excel_tem_as_mesmas_linhas_da_pagina(cliente, logado, recebimento_nfc, pagina, filtros, aba):
    logado()
    total = cliente.get(f"/paginas/{pagina}", params=filtros).json()["total"]
    r = cliente.get(f"/paginas/{pagina}/excel", params=filtros)
    assert r.status_code == 200 and r.headers["content-type"] == MIME_XLSX
    folha = load_workbook(BytesIO(r.content)).worksheets[0]
    assert folha.title == aba and folha.max_row - 1 == total > 0


# ---------------------------------------------------------------- APP-032: gráficos da Visão geral e resumo por situação
def test_visao_geral_traz_um_ponto_por_dia_e_bate_com_as_views(cliente, conexao, criar_pessoa):
    criar_pessoa(usuario="ana.teste", cargo="Operador")
    cliente.post("/login", json={"usuario": "ana.teste", "senha": "Senha-de-teste-1"})
    corpo = cliente.get("/paginas/visao-geral", params={"data_de": "2026-10-10", "data_ate": "2026-10-19"}).json()
    serie = corpo["serie"]
    assert [p["dia"] for p in serie] == [f"2026-10-{d}" for d in range(10, 20)]
    dia = "2026-10-17"
    esperado_nao_ok = (
        conexao.execute("SELECT count(*) FROM bi.contagens WHERE lido AND alarme AND data_leitura = %s", (dia,)).fetchone()[0]
        + conexao.execute("SELECT count(*) FROM bi.testes WHERE positivo AND data_resultado = %s", (dia,)).fetchone()[0]
        + conexao.execute("SELECT count(*) FROM bi.desvios WHERE data_abertura = %s", (dia,)).fetchone()[0]
    )
    esperado_registradas = (
        conexao.execute("SELECT count(*) FROM bi.contagens WHERE lido AND data_leitura = %s", (dia,)).fetchone()[0]
        + conexao.execute("SELECT count(*) FROM bi.recebimento_nfc WHERE lido AND data_leitura = %s", (dia,)).fetchone()[0]
        + conexao.execute("SELECT count(*) FROM bi.testes WHERE data_resultado = %s", (dia,)).fetchone()[0]
        + conexao.execute(
            "SELECT count(*) FROM bi.drops WHERE (concluido_em AT TIME ZONE 'America/Sao_Paulo')::date = %s", (dia,)
        ).fetchone()[0]
        + conexao.execute(
            "SELECT count(*) FROM bi.leituras_nfc WHERE (lida_em AT TIME ZONE 'America/Sao_Paulo')::date = %s", (dia,)
        ).fetchone()[0]
    )
    ponto = next(p for p in serie if p["dia"] == dia)
    assert (ponto["registradas"], ponto["nao_ok"]) == (esperado_registradas, esperado_nao_ok)
    assert sum(p["registradas"] for p in serie) > 0
    assert cliente.get("/paginas/visao-geral", params={"data_de": "2025-01-01", "data_ate": "2026-10-19"}).status_code == 422


def test_paginas_trazem_o_resumo_por_situacao_do_que_esta_filtrado(cliente, conexao, criar_pessoa):
    criar_pessoa(usuario="ana.teste", cargo="Operador")
    cliente.post("/login", json={"usuario": "ana.teste", "senha": "Senha-de-teste-1"})
    corpo = cliente.get("/paginas/recebimento").json()
    assert corpo["resumo_coluna"] == "situacao"
    por_situacao = {r["rotulo"]: r["total"] for r in corpo["resumo"]}
    assert sum(por_situacao.values()) == corpo["total"]
    esperado = dict(
        conexao.execute(
            "SELECT CASE WHEN lido AND alarme THEN 'Não ok' WHEN lido THEN 'Ok' ELSE situacao END, count(*) "
            "FROM bi.contagens WHERE origem = 'Recebimento' GROUP BY 1"
        ).fetchall()
    )
    assert por_situacao == {k: v for k, v in esperado.items()}
    # o resumo acompanha o filtro, não a página
    so_atrasados = cliente.get("/paginas/recebimento", params={"situacao": "Atrasado"}).json()
    assert [r["rotulo"] for r in so_atrasados["resumo"]] == ["Atrasado"]
    assert so_atrasados["resumo"][0]["total"] == so_atrasados["total"]
    # a ordem põe o que pede atenção primeiro
    ordem = [r["rotulo"] for r in corpo["resumo"]]
    assert ordem.index("Atrasado") < ordem.index("Ok") if {"Atrasado", "Ok"} <= set(ordem) else True
    # páginas por status (Navios, Tanques terra) também
    navios = cliente.get("/paginas/navios").json()
    assert navios["resumo_coluna"] == "status" and sum(r["total"] for r in navios["resumo"]) == navios["total"]
